"""Testing-workbook importer (Oct 2026): reads the study workbook's tabs,
matches names, converts units, resolves sides from the throwing hand,
skips flagged / text cells with a note, never overwrites."""

import io
from datetime import date

import openpyxl
import pytest

import study_workbook_spec as spec
from models import Player, Assessment, AssessmentTestType, ResearchFBOutcome
from services import study_workbook_import as swi


def _book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SH"
    ws.append([None, None, None, "ER"] + [None] * 9 + ["IR"] + [None] * 9 + ["Ratios"])
    hdr = ["First", "Last", "HAND", "THROW_Max Force (N)", "GLOVE_Max Force (N)", "Max Force_ASYM",
           "THROW_ Avg Force (N)", "GLOVE_ Avg Force (N)", "AVG FORCE_ASYM", "Max_Impulse_THROW (Ns)",
           "Max_Impulse_Glove (Ns)", "THROW_Time To Peak Force (s)", "GLOVE_Time To Peak Force (s)"]
    ws.append(hdr + hdr[3:] + ["MAX_Sh ER/IR_THROW"])
    ws.append(["Zed", "Testman", "LHP", 200, 180, 10, 190, 170, 10, 500, 480, 1.0, 1.1,
               160, 150, 6, 150, 140, 6, 450, 440, 1.2, 1.3, 1.25])
    fld = wb.create_sheet("Field Testing")
    fld.append(["FIRST", "LAST", "GRIP_THROW_kg", "Broad Jump_m", "SH ER_THROW", "HIP_IR_DRIVE", "MYSTERY"])
    fld.append(["Zed ", "Test-man", 50, 2.5, 120, 40, 7])            # close match on the name
    im = wb.create_sheet("IMTP")
    im.append([None, "First", "Last", "Peak Vertical Force [N]", "Peak Vertical Force [N] (R)"])
    im.append(["ISOT (WRONG)", "Zed", "Testman", 2000, 1000])
    st = wb.create_sheet("Strength")
    st.append(["First", "Last", "FS", "Bench Press", "HBDL", None])
    st.append(["Zed", "Testman", "355 BELTSQUAT", 200, 400, "KNEE"])
    fb = wb.create_sheet("FB VELO&SPIN")
    fb.append(["First", "Last", "AVG VELO", "PEAK VELO", "AVG SPIN RATE", "PEAK SPIN RATE", "# of Pitches", None])
    fb.append(["Zed", "Testman", 84.5, 86.0, 2100, 2300, 25, "FROM BULLPEN"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def own_db():
    """Its own database -- adding a pitcher to the shared seeded one would
    change other tests' counts."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from tests import seed
    import seed_lookups
    engine = create_engine("sqlite://")
    s = sessionmaker(bind=engine)()
    seed.build(s, engine)
    seed_lookups.seed_assessment_categories_and_tests(s)
    yield s
    s.close()


@pytest.fixture(scope="module")
def setup(own_db):
    p = Player(first_name="Zed", last_name="Testman", is_pitcher=True, active=True, throws=None)
    own_db.add(p)
    own_db.commit()
    return p


def test_plan(own_db, setup):
    db = own_db
    plan = swi.build_plan(db, _book(), date(2026, 10, 15))
    pp = plan["players"][setup.player_id]
    assert pp["throws"] == "L" and pp["set_throws"] == "L"            # from the sheet's HAND column
    v = pp["values"]
    assert abs(v["Shoulder Strength: Throwing Arm ER Peak Force"] - 200 / 4.4482216) < 0.01   # N -> lb
    assert v["Shoulder Strength: Throwing Arm ER:IR Ratio"] == pytest.approx(200 / 160, abs=1e-3)
    assert "Shoulder: Left External Rotation" in v                   # throw side of a LHP
    assert "Hip: Left Internal Rotation" in v                        # drive leg of a LHP
    assert v["Grip Strength (Seated, Throwing Hand)"] == pytest.approx(50 * 2.2046226, abs=0.01)
    assert v["Broad Jump Distance"] == pytest.approx(2.5 * 39.37008, abs=0.01)
    assert "Isometric Mid-Thigh Pull Peak Vertical Force" not in v   # flagged row skipped
    assert "Front Squat Max" not in v and v["Hex Bar Deadlift Max"] == 400
    assert any("ISOT (WRONG)" in n for n in pp["notes"]) and any("KNEE" in n for n in pp["notes"])
    assert plan["tabs"]["Field Testing"]["unknown"] == ["MYSTERY"]
    assert pp["fb"]["avg_velo"] == 84.5 and pp["fb"]["note"] == "FROM BULLPEN"
    close = [n for n in plan["names"] if n["how"] == "close"]
    assert close and close[0]["player_id"] == setup.player_id


def test_commit_never_overwrites(own_db, setup):
    db = own_db
    day = date(2026, 10, 15)
    res = swi.commit_plan(db, swi.build_plan(db, _book(), day), 1)
    assert res["values"] > 10 and res["fb_outcomes"] == 1
    assert db.query(Player).get(setup.player_id).throws == "L"
    assert db.query(ResearchFBOutcome).filter_by(player_id=setup.player_id).one().n_pitches == 25
    notes = " ".join(a.notes or "" for a in db.query(Assessment).filter_by(player_id=setup.player_id))
    assert "ISOT (WRONG)" in notes
    again = swi.build_plan(db, _book(), day)
    assert again["totals"]["new_values"] == 0 and again["totals"]["already_recorded"] > 10


def test_override_skips_a_name(own_db, setup):
    db = own_db
    plan = swi.build_plan(db, _book(), date(2026, 10, 16), {"zed|test-man": 0})
    n = [x for x in plan["names"] if x["key"] == "zed|test-man"][0]
    assert n["player_id"] is None and n["how"] == "skipped"


def test_new_test_types_cover_every_column():
    names = {n for _c, n, _u in spec.new_test_types()}
    assert "IMTP RFD 0-100 ms" in names and "Shoulder ER Impulse (Throwing Arm)" in names
    assert "Countermovement Jump Height" not in names                # already a GBO test
